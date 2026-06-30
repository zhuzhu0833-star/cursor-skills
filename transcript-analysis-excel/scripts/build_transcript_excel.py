#!/usr/bin/env python3
"""Build a styled multi-sheet Excel workbook from transcript JSON."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any

GPA_GRADES = {
    "A+", "A", "A-", "B+", "B", "B-", "C+", "C", "C-", "D+", "D", "D-", "F"
}

GRADE_SORT = {
    "A+": 4.0, "A": 4.0, "A-": 3.7,
    "B+": 3.3, "B": 3.0, "B-": 2.7,
    "C+": 2.3, "C": 2.0, "C-": 1.7,
    "D+": 1.3, "D": 1.0, "D-": 0.7, "F": 0.0,
}

SHEET_COLUMNS = {
    "学生信息": ["Field", "Value"],
    "学业轨迹": ["Period", "School/Program", "Major/Plan", "Action"],
    "课程明细": [
        "Term", "Program", "Plan", "End of Study",
        "Course Code", "Course Description", "Credits", "Grade", "Course Topic",
    ],
    "学期GPA": [
        "Term", "Program", "Plan", "Term GPA", "Term Credits",
        "Cum GPA", "Cum Credits", "Notes",
    ],
    "成绩分布": ["Grade", "Count"],
    "里程碑": [
        "Program", "Milestone", "Level", "Status",
        "Date Completed", "Attempt Status", "Attempted On",
    ],
}


def load_data(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as f:
        data = json.load(f)
    for key in ("student_info", "academic_history", "courses", "term_gpa", "milestones"):
        data.setdefault(key, {} if key == "student_info" else [])
    return data


def student_info_rows(info: dict[str, Any]) -> list[dict[str, str]]:
    mapping = [
        ("University", "university"),
        ("Student Name", "student_name"),
        ("Student ID", "student_id"),
        ("Date Printed", "date_printed"),
        ("Document Type", "document_type"),
        ("Current Program", "current_program"),
        ("Current Major", "current_major"),
        ("Concentrations", "concentrations"),
        ("Cumulative GPA", "cumulative_gpa"),
        ("Cumulative Credits Earned", "cumulative_credits"),
    ]
    rows = []
    for label, key in mapping:
        val = info.get(key, "")
        if val not in ("", None):
            rows.append({"Field": label, "Value": str(val)})
    return rows


def history_rows(items: list[dict[str, Any]]) -> list[dict[str, str]]:
    return [
        {
            "Period": item.get("period", ""),
            "School/Program": item.get("school_program", item.get("school", "")),
            "Major/Plan": item.get("major_plan", item.get("major", "")),
            "Action": item.get("action", ""),
        }
        for item in items
    ]


def course_rows(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for item in items:
        rows.append({
            "Term": item.get("term", ""),
            "Program": item.get("program", ""),
            "Plan": item.get("plan", ""),
            "End of Study": item.get("end_of_study", ""),
            "Course Code": item.get("course_code", ""),
            "Course Description": item.get("course_description", ""),
            "Credits": item.get("credits", ""),
            "Grade": item.get("grade", ""),
            "Course Topic": item.get("course_topic", ""),
        })
    return rows


def term_gpa_rows(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "Term": item.get("term", ""),
            "Program": item.get("program", ""),
            "Plan": item.get("plan", ""),
            "Term GPA": item.get("term_gpa", ""),
            "Term Credits": item.get("term_credits", ""),
            "Cum GPA": item.get("cum_gpa", ""),
            "Cum Credits": item.get("cum_credits", ""),
            "Notes": item.get("notes", ""),
        }
        for item in items
    ]


def milestone_rows(items: list[dict[str, Any]]) -> list[dict[str, str]]:
    return [
        {
            "Program": item.get("program", ""),
            "Milestone": item.get("milestone", ""),
            "Level": item.get("level", ""),
            "Status": item.get("status", ""),
            "Date Completed": item.get("date_completed", ""),
            "Attempt Status": item.get("attempt_status", ""),
            "Attempted On": item.get("attempted_on", ""),
        }
        for item in items
    ]


def grade_distribution(courses: list[dict[str, Any]]) -> list[dict[str, Any]]:
    counts: dict[str, int] = {}
    for course in courses:
        grade = str(course.get("grade", "")).strip()
        if grade in GPA_GRADES:
            counts[grade] = counts.get(grade, 0) + 1
    ordered = sorted(counts.keys(), key=lambda g: GRADE_SORT.get(g, -1), reverse=True)
    return [{"Grade": g, "Count": counts[g]} for g in ordered]


def default_output_name(data: dict[str, Any]) -> str:
    info = data.get("student_info", {})
    name = info.get("student_name", "Student")
    university = info.get("university", "Transcript")
    parts = name.split()
    if len(parts) >= 2:
        given = "_".join(parts[:-1])
        family = parts[-1]
        name_slug = f"{given}_{family}"
    else:
        name_slug = re.sub(r"[^\w]+", "_", name)
    inst_slug = re.sub(r"[^\w]+", "_", university.split()[0])
    return f"{name_slug}_{inst_slug}_Transcript.xlsx"


def style_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = thin_border
        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value), default=0)
            letter = get_column_letter(col[0].column)
            ws.column_dimensions[letter].width = min(max_len + 4, 50)
        ws.freeze_panes = "A2"

    wb.save(path)


def build_excel(data: dict[str, Any], output_path: Path) -> None:
    import pandas as pd

    sheets = {
        "学生信息": student_info_rows(data["student_info"]),
        "学业轨迹": history_rows(data["academic_history"]),
        "课程明细": course_rows(data["courses"]),
        "学期GPA": term_gpa_rows(data["term_gpa"]),
        "成绩分布": grade_distribution(data["courses"]),
        "里程碑": milestone_rows(data["milestones"]),
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, rows in sheets.items():
            columns = SHEET_COLUMNS[sheet_name]
            df = pd.DataFrame(rows, columns=columns) if rows else pd.DataFrame(columns=columns)
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    style_workbook(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="Path to transcript JSON")
    parser.add_argument("--output-dir", required=True, help="Directory for output .xlsx")
    parser.add_argument("--output-name", help="Override output filename")
    args = parser.parse_args()

    input_path = Path(args.input).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()

    if not input_path.is_file():
        print(f"Error: input file not found: {input_path}", file=sys.stderr)
        return 1

    data = load_data(input_path)
    output_name = args.output_name or default_output_name(data)
    output_path = output_dir / output_name

    try:
        build_excel(data, output_path)
    except Exception as exc:
        print(f"Error building Excel: {exc}", file=sys.stderr)
        return 1

    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
