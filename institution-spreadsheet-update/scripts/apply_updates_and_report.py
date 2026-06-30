#!/usr/bin/env python3
"""Apply institution spreadsheet updates, highlight changes in red, export PDF report."""

from __future__ import annotations

import argparse
import csv
import json
import sys
from datetime import date
from pathlib import Path
from typing import Any

EMPTY_VALUES = {"", "—", "-", "n/a", "na", "null", "none"}


def normalize(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in EMPTY_VALUES:
        return ""
    return " ".join(text.split())


def load_table(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [normalize(h) for h in next(rows_iter)]
        rows: list[dict[str, str]] = []
        for row in rows_iter:
            record = {
                headers[i]: normalize(row[i]) if i < len(row) and row[i] is not None else ""
                for i in range(len(headers))
            }
            if any(record.values()):
                rows.append(record)
        wb.close()
        return headers, rows

    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        headers = list(reader.fieldnames or [])
        rows = [{h: normalize(row.get(h, "")) for h in headers} for row in reader]
    return headers, rows


def load_changes(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, list):
        return {"checked_at": date.today().isoformat(), "schools": data, "unverified": [], "summary_zh": []}
    data.setdefault("schools", [])
    data.setdefault("unverified", [])
    data.setdefault("summary_zh", [])
    data.setdefault("checked_at", date.today().isoformat())
    return data


def find_row_index(rows: list[dict[str, str]], school: dict[str, Any], key_fields: list[str]) -> int | None:
    for idx, row in enumerate(rows):
        for field in key_fields:
            val = school.get(field.replace("_col", "")) or school.get(field)
            if val and normalize(row.get(field, "")) == normalize(val):
                return idx
    # fallback: match school_id or names inside row
    for idx, row in enumerate(rows):
        if school.get("school_id") and normalize(row.get("school_id", "")) == normalize(school["school_id"]):
            return idx
        if school.get("school_name_en") and normalize(row.get("school_name_en", "")) == normalize(school["school_name_en"]):
            return idx
        if school.get("school_name_zh") and normalize(row.get("school_name_zh", "")) == normalize(school["school_name_zh"]):
            return idx
    return None


def apply_changes(
    headers: list[str],
    rows: list[dict[str, str]],
    schools: list[dict[str, Any]],
    checked_at: str,
) -> tuple[list[dict[str, str]], list[tuple[int, str]], list[dict[str, Any]]]:
    updated = [row.copy() for row in rows]
    changed_cells: list[tuple[int, str]] = []
    applied_schools: list[dict[str, Any]] = []

    for school in schools:
        idx = find_row_index(updated, school, ["school_id", "school_name_en", "school_name_zh"])
        if idx is None:
            print(f"Warning: school not found in table: {school.get('school_id') or school.get('school_name_en')}", file=sys.stderr)
            continue

        applied_changes = []
        for change in school.get("changes", []):
            field = change["field"]
            if field not in headers:
                print(f"Warning: unknown column '{field}' — skipped", file=sys.stderr)
                continue
            new_val = normalize(change.get("new", ""))
            old_val = normalize(updated[idx].get(field, ""))
            if normalize(change.get("old", old_val)) != old_val and old_val != new_val:
                # keep going; trust explicit new value
                pass
            if old_val == new_val:
                continue
            updated[idx][field] = new_val
            changed_cells.append((idx, field))
            applied_changes.append({**change, "old": old_val, "new": new_val})

        if applied_changes:
            if "last_updated" in headers:
                updated[idx]["last_updated"] = checked_at
            applied_schools.append({**school, "changes": applied_changes})

    return updated, changed_cells, applied_schools


def write_csv(path: Path, headers: list[str], rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_highlighted_xlsx(
    path: Path,
    headers: list[str],
    rows: list[dict[str, str]],
    changed_cells: list[tuple[int, str]],
) -> None:
    import openpyxl
    from openpyxl.styles import Font

    red = Font(color="FF0000")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Updated"

    col_index = {h: i + 1 for i, h in enumerate(headers)}
    changed_set = set(changed_cells)

    for c, header in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=header)

    for r, row in enumerate(rows, start=2):
        data_idx = r - 2
        for header in headers:
            c = col_index[header]
            value = row.get(header, "")
            cell = ws.cell(row=r, column=c, value=value)
            if (data_idx, header) in changed_set:
                cell.font = red

    wb.save(path)


def find_cjk_font() -> Path | None:
    candidates = [
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
        Path("/System/Library/Fonts/STHeiti Light.ttc"),
        Path("/System/Library/Fonts/PingFang.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return path
    return None


def register_cjk_font(pdf: Any, font_path: Path) -> None:
    """Register a Unicode font for Chinese PDF output."""
    try:
        pdf.add_font("CJK", "", str(font_path))
    except Exception:
        # Some fpdf2 versions need uni=True for TTF/TTC
        pdf.add_font("CJK", "", str(font_path), uni=True)


def pdf_write_lines(pdf: Any, lines: str, line_height: float = 7) -> None:
    width = pdf.epw
    for paragraph in lines.split("\n"):
        pdf.multi_cell(width, line_height, paragraph)
        pdf.ln(1)


def pdf_write_bullets(pdf: Any, items: list[str], line_height: float = 6) -> None:
    width = pdf.epw
    for item in items:
        pdf.multi_cell(width, line_height, f"- {item}")
        pdf.ln(1)


def write_pdf_report(
    path: Path,
    input_path: Path,
    applied_schools: list[dict[str, Any]],
    checked_at: str,
    unverified: list[dict[str, Any]],
    summary_zh: list[str],
    font_path: Path | None,
) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)

    resolved_font = font_path or find_cjk_font()
    if resolved_font is None:
        raise RuntimeError("No CJK font found. Pass --font /path/to/font.ttc")

    register_cjk_font(pdf, resolved_font)
    pdf.add_page()
    pdf.set_font("CJK", size=14)
    pdf.cell(pdf.epw, 10, "院校表格更新报告", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(3)

    total_changes = sum(len(s.get("changes", [])) for s in applied_schools)
    pdf.set_font("CJK", size=11)
    pdf_write_lines(
        pdf,
        f"检查日期：{checked_at}\n"
        f"源文件：{input_path.name}\n"
        f"变更院校数：{len(applied_schools)}\n"
        f"变更字段数：{total_changes}",
    )
    pdf.ln(2)

    if summary_zh:
        pdf.set_font("CJK", size=12)
        pdf.cell(pdf.epw, 8, "变更解读", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("CJK", size=10)
        pdf_write_bullets(pdf, summary_zh)
        pdf.ln(2)

    pdf.set_font("CJK", size=12)
    pdf.cell(pdf.epw, 8, "变更详情", new_x="LMARGIN", new_y="NEXT")

    col_w = pdf.epw / 3
    for school in applied_schools:
        name_en = school.get("school_name_en", "")
        name_zh = school.get("school_name_zh", "")
        title = f"{name_en} ({name_zh})" if name_zh else name_en
        pdf.set_font("CJK", size=11)
        pdf.multi_cell(pdf.epw, 7, title)
        pdf.ln(1)
        pdf.set_font("CJK", size=9)

        pdf.set_fill_color(240, 240, 240)
        pdf.cell(col_w, 7, "字段", border=1, fill=True)
        pdf.cell(col_w, 7, "旧值", border=1, fill=True)
        pdf.cell(col_w, 7, "新值", border=1, fill=True, new_x="LMARGIN", new_y="NEXT")

        for change in school.get("changes", []):
            label = change.get("label_zh") or change.get("field", "")
            old = (change.get("old", "") or "-")[:120]
            new = (change.get("new", "") or "-")[:120]
            pdf.cell(col_w, 7, label, border=1)
            pdf.cell(col_w, 7, old, border=1)
            pdf.cell(col_w, 7, new, border=1, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)

    if unverified:
        pdf.set_font("CJK", size=12)
        pdf.cell(pdf.epw, 8, "未能核实", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("CJK", size=9)
        pdf_write_bullets(
            pdf,
            [f"{item.get('school_id', '')} / {item.get('field', '')}: {item.get('reason', '')}" for item in unverified],
        )

    pdf.output(str(path))


def load_summary_md(path: Path | None) -> list[str]:
    if path is None or not path.exists():
        return []
    lines = []
    for line in path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith(("#", "-", "*", "•")):
            stripped = stripped.lstrip("#-*• ").strip()
        if stripped:
            lines.append(stripped)
    return lines


def main() -> int:
    parser = argparse.ArgumentParser(description="Apply spreadsheet updates with red highlights and PDF report.")
    parser.add_argument("--input", required=True, type=Path, help="Source CSV or XLSX")
    parser.add_argument("--changes", required=True, type=Path, help="Changes JSON file")
    parser.add_argument("--output-dir", type=Path, default=None, help="Output directory (default: input folder)")
    parser.add_argument("--date", default=date.today().strftime("%Y%m%d"), help="Date stamp YYYYMMDD")
    parser.add_argument("--font", type=Path, default=None, help="TTF/TTC font for PDF CJK text")
    parser.add_argument("--summary-md", type=Path, default=None, help="Optional markdown bullets for 变更解读")
    args = parser.parse_args()

    input_path = args.input.expanduser().resolve()
    changes_path = args.changes.expanduser().resolve()
    output_dir = (args.output_dir or input_path.parent).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    payload = load_changes(changes_path)
    headers, rows = load_table(input_path)
    checked_at = payload.get("checked_at", date.today().isoformat())

    updated_rows, changed_cells, applied_schools = apply_changes(
        headers, rows, payload.get("schools", []), checked_at
    )

    stem = input_path.stem
    date_stamp = args.date

    csv_out = output_dir / f"{stem}_updated_{date_stamp}.csv"
    xlsx_out = output_dir / f"{stem}_updated_{date_stamp}.xlsx"
    pdf_out = output_dir / f"{stem}_update_report_{date_stamp}.pdf"
    json_out = output_dir / f"{stem}_changes_{date_stamp}.json"

    write_csv(csv_out, headers, updated_rows)
    write_highlighted_xlsx(xlsx_out, headers, updated_rows, changed_cells)

    summary = list(payload.get("summary_zh") or [])
    summary.extend(load_summary_md(args.summary_md))

    archive = {
        "checked_at": checked_at,
        "input": str(input_path),
        "schools": applied_schools,
        "unverified": payload.get("unverified", []),
        "summary_zh": summary,
        "outputs": {
            "csv": str(csv_out),
            "xlsx": str(xlsx_out),
            "pdf": str(pdf_out),
        },
    }
    json_out.write_text(json.dumps(archive, ensure_ascii=False, indent=2), encoding="utf-8")

    write_pdf_report(
        pdf_out,
        input_path,
        applied_schools,
        checked_at,
        payload.get("unverified", []),
        summary,
        args.font,
    )

    manifest = {
        "generated_at": checked_at,
        "input": str(input_path),
        "csv": str(csv_out),
        "xlsx": str(xlsx_out),
        "pdf": str(pdf_out),
        "changes_json": str(json_out),
        "school_count": len(rows),
        "change_count": len(applied_schools),
        "field_change_count": sum(len(s.get("changes", [])) for s in applied_schools),
    }
    manifest_path = output_dir / f"{stem}_report_manifest_{date_stamp}.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
