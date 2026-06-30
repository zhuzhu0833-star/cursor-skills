#!/usr/bin/env python3
"""Apply updates to merged-header institution XLSX (rows 1-2 headers, data from row 3)."""

from __future__ import annotations

import argparse
import csv
import json
import sys
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any

# Reuse helpers from sibling script
sys.path.insert(0, str(Path(__file__).resolve().parent))
from apply_updates_and_report import (  # noqa: E402
    find_cjk_font,
    load_changes,
    normalize,
    register_cjk_font,
    write_pdf_report,
)

HEADER_ROW = 1
SUBHEADER_ROW = 2
DATA_START_ROW = 3


def build_column_map(ws) -> dict[str, int]:
    """Map composite field names to 1-based column indices."""
    col_map: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        h1 = normalize(ws.cell(HEADER_ROW, c).value)
        h2 = normalize(ws.cell(SUBHEADER_ROW, c).value)
        if h1 and h2:
            col_map[f"{h1}|{h2}"] = c
        if h1:
            col_map[h1] = c
        if h2:
            col_map[h2] = c
    return col_map


def cell_value_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return normalize(value)


def read_data_rows(ws, col_map: dict[str, int]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for r in range(DATA_START_ROW, ws.max_row + 1):
        name = ws.cell(r, col_map.get("学校名称", 2)).value
        if not name:
            continue
        record: dict[str, str] = {"_row": str(r)}
        for field, c in col_map.items():
            if "|" in field or field in {
                "USNEWS综合排名排名",
                "学校名称",
                "录取要求链接",
                "School Name",
                "本科人数",
                "参考学费",
                "平均GPA",
                "雅思",
                "Duolingo",
                "录取率",
                "学校描述",
            }:
                val = ws.cell(r, c).value
                if val is not None:
                    record[field] = cell_value_str(val)
        rows.append(record)
    return rows


def find_row(ws, col_map: dict[str, int], school: dict[str, Any]) -> int | None:
    name_col = col_map.get("学校名称", 2)
    en_col = col_map.get("School Name", 4)
    for r in range(DATA_START_ROW, ws.max_row + 1):
        zh = normalize(ws.cell(r, name_col).value)
        en = normalize(ws.cell(r, en_col).value)
        if school.get("school_name_zh") and zh == normalize(school["school_name_zh"]):
            return r
        if school.get("school_name_en") and en == normalize(school["school_name_en"]):
            return r
    return None


def apply_to_workbook(
    ws,
    col_map: dict[str, int],
    schools: list[dict[str, Any]],
    checked_at: str,
) -> tuple[list[tuple[int, int]], list[dict[str, Any]]]:
    changed_cells: list[tuple[int, int]] = []
    applied_schools: list[dict[str, Any]] = []

    for school in schools:
        row = find_row(ws, col_map, school)
        if row is None:
            print(
                f"Warning: school not found: {school.get('school_name_zh') or school.get('school_name_en')}",
                file=sys.stderr,
            )
            continue

        applied_changes = []
        for change in school.get("changes", []):
            field = change["field"]
            col = col_map.get(field)
            if col is None:
                print(f"Warning: unknown column '{field}' — skipped", file=sys.stderr)
                continue
            old_val = cell_value_str(ws.cell(row, col).value)
            new_val = normalize(change.get("new", ""))
            if normalize(old_val) == new_val:
                continue
            ws.cell(row, col, value=new_val if new_val else None)
            changed_cells.append((row, col))
            applied_changes.append({**change, "old": old_val, "new": new_val})

        if applied_changes:
            applied_schools.append({**school, "changes": applied_changes})

    return changed_cells, applied_schools


def highlight_changes(ws, changed_cells: list[tuple[int, int]], original_fonts: dict) -> None:
    from openpyxl.styles import Font

    red = Font(color="FF0000")
    for row, col in changed_cells:
        cell = ws.cell(row, col)
        base = original_fonts.get((row, col))
        if base:
            cell.font = copy(base)
            cell.font = Font(
                name=cell.font.name,
                size=cell.font.size,
                bold=cell.font.bold,
                italic=cell.font.italic,
                color="FF0000",
            )
        else:
            cell.font = red


def export_flat_csv(path: Path, ws, col_map: dict[str, int]) -> None:
    fields = sorted({f for f in col_map if "|" in f or f in {
        "USNEWS综合排名排名", "学校名称", "录取要求链接", "School Name",
        "本科人数", "参考学费", "平均GPA", "雅思", "Duolingo", "录取率",
    }})
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for r in range(DATA_START_ROW, ws.max_row + 1):
            if not ws.cell(r, col_map.get("学校名称", 2)).value:
                continue
            row_dict = {}
            for field in fields:
                c = col_map[field]
                row_dict[field] = cell_value_str(ws.cell(r, c).value)
            writer.writerow(row_dict)


def main() -> int:
    parser = argparse.ArgumentParser(description="Apply updates to merged-header institution XLSX.")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--changes", required=True, type=Path)
    parser.add_argument("--output-dir", type=Path, default=None)
    parser.add_argument("--date", default=date.today().strftime("%Y%m%d"))
    parser.add_argument("--font", type=Path, default=None)
    parser.add_argument("--summary-md", type=Path, default=None)
    args = parser.parse_args()

    import openpyxl

    input_path = args.input.expanduser().resolve()
    output_dir = (args.output_dir or input_path.parent).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    payload = load_changes(args.changes.expanduser().resolve())
    checked_at = payload.get("checked_at", date.today().isoformat())

    wb = openpyxl.load_workbook(input_path)
    ws = wb.active
    col_map = build_column_map(ws)

    original_fonts = {}
    for r in range(DATA_START_ROW, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            if cell.font:
                original_fonts[(r, c)] = copy(cell.font)

    changed_cells, applied_schools = apply_to_workbook(
        ws, col_map, payload.get("schools", []), checked_at
    )
    highlight_changes(ws, changed_cells, original_fonts)

    stem = input_path.stem
    date_stamp = args.date
    xlsx_out = output_dir / f"{stem}_updated_{date_stamp}.xlsx"
    csv_out = output_dir / f"{stem}_updated_{date_stamp}.csv"
    pdf_out = output_dir / f"{stem}_update_report_{date_stamp}.pdf"
    json_out = output_dir / f"{stem}_changes_{date_stamp}.json"

    wb.save(xlsx_out)
    export_flat_csv(csv_out, ws, col_map)
    wb.close()

    summary = list(payload.get("summary_zh") or [])
    if args.summary_md and args.summary_md.exists():
        for line in args.summary_md.read_text(encoding="utf-8").splitlines():
            s = line.strip().lstrip("#-*• ").strip()
            if s:
                summary.append(s)

    archive = {
        "checked_at": checked_at,
        "input": str(input_path),
        "schools": applied_schools,
        "unverified": payload.get("unverified", []),
        "summary_zh": summary,
        "outputs": {"csv": str(csv_out), "xlsx": str(xlsx_out), "pdf": str(pdf_out)},
    }
    json_out.write_text(json.dumps(archive, ensure_ascii=False, indent=2), encoding="utf-8")

    write_pdf_report(
        pdf_out,
        input_path,
        applied_schools,
        checked_at,
        payload.get("unverified", []),
        summary,
        args.font,
    )

    manifest = {
        "generated_at": checked_at,
        "input": str(input_path),
        "csv": str(csv_out),
        "xlsx": str(xlsx_out),
        "pdf": str(pdf_out),
        "changes_json": str(json_out),
        "change_count": len(applied_schools),
        "field_change_count": sum(len(s.get("changes", [])) for s in applied_schools),
    }
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
