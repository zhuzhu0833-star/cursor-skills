#!/usr/bin/env python3
"""Generate WeChat-friendly PDF/PNG reports and highlighted Excel exports."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

SKILL_ROOT = Path(__file__).resolve().parent.parent
FONT_PATH = Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf")
YELLOW = "#FFF59D"
YELLOW_FILL = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")
HEADER_FILL = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
IMG_WIDTH = 1080
IMG_MAX_HEIGHT = 3600
MARGIN = 48

DISPLAY_FIELDS = [
    ("school_name_zh", "中文校名"),
    ("school_name_en", "English Name"),
    ("test_policy", "标化政策"),
    ("rd_deadline", "RD 截止"),
    ("ea_deadline", "EA 截止"),
    ("ed_deadline", "ED 截止"),
    ("toefl_min_ibt", "TOEFL"),
    ("ielts_min", "IELTS"),
    ("english_waiver_policy", "英语豁免"),
    ("need_blind_international", "国际生 Need-Blind"),
    ("intl_scholarship_available", "国际生奖学金"),
    ("application_fee", "申请费"),
    ("tuition_out_of_state", "州外学费"),
    ("notes_zh", "备注"),
]

FIELD_KEY = {k: v for k, v in DISPLAY_FIELDS}


def _register_font() -> str:
    name = "ArialUnicode"
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(name, str(FONT_PATH)))
    return name


def _pil_font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_PATH), size)


def _wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    if not text:
        return ["—"]
    words = list(text)
    lines: list[str] = []
    current = ""
    for ch in words:
        trial = current + ch
        if draw.textlength(trial, font=font) <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines or ["—"]


def _changed_fields_map(changes: list[dict[str, Any]]) -> dict[str, set[str]]:
    result: dict[str, set[str]] = {}
    for item in changes:
        sid = item["school_id"]
        result[sid] = {c["field"] for c in item["changes"]}
    return result


def export_highlighted_xlsx(
    snapshot: dict[str, Any],
    changes: list[dict[str, Any]],
    output_path: Path,
) -> Path:
    changed_map = _changed_fields_map(changes)
    wb = Workbook()
    ws = wb.active
    ws.title = "政策监控"

    headers = ["school_id"] + [label for _, label in DISPLAY_FIELDS[1:]]
    policy_keys = ["school_name_en", "school_name_zh"] + [k for k, _ in DISPLAY_FIELDS[2:]]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_idx = 2
    schools = snapshot.get("schools", {})
    for sid in sorted(schools.keys(), key=lambda s: schools[s].get("school_name_zh") or schools[s].get("school_name_en", s)):
        entry = schools[sid]
        policy = entry.get("policy", {})
        changed = changed_map.get(sid, set())

        values = [sid]
        keys = ["school_id"] + policy_keys[1:]
        for key in keys[1:]:
            if key in ("school_name_en", "school_name_zh"):
                values.append(entry.get(key, ""))
            else:
                values.append(policy.get(key, ""))

        for col, (key, val) in enumerate(zip(keys, values), 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            field_key = key if key != "school_id" else None
            if field_key and field_key in changed:
                cell.fill = YELLOW_FILL
        row_idx += 1

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18 if col > 2 else 14
    ws.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def _build_summary_lines(snapshot: dict[str, Any], previous: dict[str, Any] | None, changes: list[dict[str, Any]]) -> list[str]:
    checked = snapshot.get("checked_at", date.today().isoformat())
    prev = (previous or {}).get("checked_at", "无")
    return [
        f"检查日期：{checked}",
        f"上次快照：{prev}",
        f"监控院校：{len(snapshot.get('schools', {}))} 所（全部表格学校）",
        f"本次变更：{len(changes)} 所",
        "黄色高亮 = 本次更新字段",
    ]


def export_png_pages(
    snapshot: dict[str, Any],
    previous: dict[str, Any] | None,
    changes: list[dict[str, Any]],
    output_dir: Path,
    stamp: str,
) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    title_font = _pil_font(42)
    h_font = _pil_font(30)
    body_font = _pil_font(24)

    summary = _build_summary_lines(snapshot, previous, changes)
    rendered_paths: list[Path] = []
    page_num = 0

    def render_page(lines: list[str]) -> Path:
        nonlocal page_num
        img_height = min(
            max(_estimate_height(lines, title_font, h_font, body_font) + 2 * MARGIN, 800),
            IMG_MAX_HEIGHT,
        )
        img = Image.new("RGB", (IMG_WIDTH, img_height), "white")
        draw = ImageDraw.Draw(img)
        y = MARGIN
        for line in lines:
            if line.startswith("美国本科"):
                font = title_font
            elif line.startswith("【"):
                font = h_font
            else:
                font = body_font
            for ln in _wrap_text(draw, line, font, IMG_WIDTH - 2 * MARGIN):
                is_changed = "← 更新" in ln
                if is_changed:
                    tw = draw.textlength(ln, font=font)
                    draw.rectangle([MARGIN - 4, y - 2, MARGIN + tw + 8, y + 28], fill=YELLOW)
                color = "#1565C0" if font == title_font else "#0D47A1" if font == h_font else "#333333"
                draw.text((MARGIN, y), ln, fill=color, font=font)
                y += 34 if font == title_font else 30 if font == h_font else 28
        page_num += 1
        out = output_dir / f"policy_update_{stamp}_p{page_num}.png"
        img.save(out, format="PNG", optimize=True)
        rendered_paths.append(out)
        return out

    if not changes:
        render_page(["美国本科院校政策更新", *summary, "", "✅ 与上次相比未发现政策变更", "", "数据来源：各校官网 / Common Data Set"])
        return rendered_paths

    intro = ["美国本科院校政策更新", *summary, ""]
    school_blocks: list[list[str]] = []
    for item in changes:
        name = item.get("school_name_zh") or item.get("school_name_en", "")
        en = item.get("school_name_en", "")
        header = f"【{name}】{en}" if name != en else f"【{en}】"
        block = [header]
        for ch in item["changes"]:
            label = ch.get("label_zh", ch["field"])
            old = ch.get("old") or "—"
            new = ch.get("new") or "—"
            block.append(f"  • {label}")
            block.append(f"    旧：{old}")
            block.append(f"    新：{new}  ← 更新")
        block.append("")
        school_blocks.append(block)

    current_lines = list(intro)
    for block in school_blocks:
        trial = current_lines + block
        if _estimate_height(trial, title_font, h_font, body_font) + 2 * MARGIN > IMG_MAX_HEIGHT and len(current_lines) > len(intro):
            render_page(current_lines)
            current_lines = ["美国本科院校政策更新（续）", ""] + block
        else:
            current_lines.extend(block)

    if current_lines:
        render_page(current_lines)

    return rendered_paths


def _estimate_height(
    lines: list[str],
    title_font: ImageFont.FreeTypeFont,
    h_font: ImageFont.FreeTypeFont,
    body_font: ImageFont.FreeTypeFont,
) -> int:
    dummy = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    total = 0
    for line in lines:
        if line.startswith("美国本科"):
            font = title_font
            lh = 34
        elif line.startswith("【"):
            font = h_font
            lh = 30
        else:
            font = body_font
            lh = 28
        wrapped = _wrap_text(dummy, line, font, IMG_WIDTH - 2 * MARGIN)
        total += lh * max(1, len(wrapped))
    return total


def export_pdf(
    snapshot: dict[str, Any],
    previous: dict[str, Any] | None,
    changes: list[dict[str, Any]],
    output_path: Path,
) -> Path:
    font_name = _register_font()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleCN", parent=styles["Title"], fontName=font_name, fontSize=18, leading=24)
    body_style = ParagraphStyle("BodyCN", parent=styles["Normal"], fontName=font_name, fontSize=10, leading=14)
    changed_style = ParagraphStyle(
        "ChangedCN",
        parent=body_style,
        backColor=colors.HexColor(YELLOW),
        fontName=font_name,
    )

    doc = SimpleDocTemplate(str(output_path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
    story: list[Any] = []

    story.append(Paragraph("美国本科院校政策更新报告", title_style))
    story.append(Spacer(1, 8))
    for line in _build_summary_lines(snapshot, previous, changes):
        story.append(Paragraph(line, body_style))
    story.append(Spacer(1, 12))

    if not changes:
        story.append(Paragraph("✅ 与上次快照相比，未发现政策字段变更。", body_style))
    else:
        table_data = [["院校", "字段", "旧值", "新值"]]
        highlight_rows: list[int] = []
        row_i = 1
        for item in changes:
            name = item.get("school_name_zh") or item.get("school_name_en", "")
            en = item.get("school_name_en", "")
            display_name = f"{name}<br/>{en}" if name and name != en else en
            for ch in item["changes"]:
                table_data.append([
                    display_name if ch == item["changes"][0] else "",
                    ch.get("label_zh", ch["field"]),
                    ch.get("old") or "—",
                    ch.get("new") or "—",
                ])
                highlight_rows.append(row_i)
                row_i += 1

        table = Table(table_data, colWidths=[55 * mm, 35 * mm, 45 * mm, 45 * mm], repeatRows=1)
        style_commands: list[Any] = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1565C0")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
        ]
        for r in highlight_rows:
            style_commands.append(("BACKGROUND", (3, r), (3, r), colors.HexColor(YELLOW)))
        table.setStyle(TableStyle(style_commands))
        story.append(table)

    story.append(Spacer(1, 10))
    story.append(Paragraph("说明：黄色高亮为本次更新字段，适合转发微信朋友圈。", body_style))
    doc.build(story)
    return output_path


def generate_all_reports(
    snapshot: dict[str, Any],
    previous: dict[str, Any] | None,
    changes: list[dict[str, Any]],
    report_dir: Path,
) -> dict[str, Any]:
    stamp = date.today().strftime("%Y%m%d")
    report_dir.mkdir(parents=True, exist_ok=True)

    xlsx_path = export_highlighted_xlsx(snapshot, changes, report_dir / f"colleges_highlighted_{stamp}.xlsx")
    pdf_path = export_pdf(snapshot, previous, changes, report_dir / f"policy_update_{stamp}.pdf")
    png_paths = export_png_pages(snapshot, previous, changes, report_dir, stamp)

    manifest = {
        "generated_at": date.today().isoformat(),
        "xlsx": str(xlsx_path),
        "pdf": str(pdf_path),
        "png": [str(p) for p in png_paths],
        "school_count": len(snapshot.get("schools", {})),
        "change_count": len(changes),
    }
    manifest_path = report_dir / f"report_manifest_{stamp}.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest["manifest"] = str(manifest_path)
    return manifest


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("--snapshot", type=Path, required=True)
    parser.add_argument("--previous", type=Path)
    parser.add_argument("--changes", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    snapshot = json.loads(args.snapshot.read_text(encoding="utf-8"))
    previous = json.loads(args.previous.read_text(encoding="utf-8")) if args.previous and args.previous.exists() else None
    changes = json.loads(args.changes.read_text(encoding="utf-8"))
    result = generate_all_reports(snapshot, previous, changes, args.output_dir)
    print(json.dumps(result, ensure_ascii=False, indent=2))
